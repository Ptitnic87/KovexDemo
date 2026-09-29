# Fichier : src/infrastructure/classeur_modele.py
"""Export du modèle de rôles au format Excel.

Trois feuilles plates plutôt qu'une feuille dense : chacune est un tableau
filtrable et pivotable dans Excel. Concaténer droits et membres dans une
cellule aurait produit un fichier lisible d'un coup d'œil mais inexploitable —
et une cellule Excel est plafonnée à 32 767 caractères, ce qu'un rôle à
plusieurs milliers de membres dépasse.

Une quatrième feuille énonce les filtres appliqués. Un extrait qui ne dit pas
ce qu'il a filtré ne permet pas de savoir si un rôle absent a été écarté ou
n'existe pas.

Ce module ne connaît aucun libellé : `traduire(cle, params)` lui est fourni,
conformément à la règle du produit.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Callable, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.infrastructure import palette
from src.infrastructure.branding import NOM_PRODUIT

#: Largeur maximale d'une colonne : au-delà, un identifiant long étalerait la
#: feuille sur plusieurs écrans.
LARGEUR_MAX = 60


def _entete(feuille, colonnes: List[str]) -> None:
    feuille.append(colonnes)
    fond = PatternFill("solid", start_color=palette.BANDEAU.lstrip("#"))
    police = Font(bold=True, color=palette.SUR_BANDEAU.lstrip("#"))
    for cellule in feuille[1]:
        cellule.fill = fond
        cellule.font = police
        cellule.alignment = Alignment(vertical="center")
    feuille.freeze_panes = "A2"


def _ajuster(feuille) -> None:
    """Largeur des colonnes d'après leur contenu réel."""
    for index, colonne in enumerate(feuille.columns, start=1):
        largeur = max((len(str(c.value)) for c in colonne if c.value is not None),
                      default=0)
        feuille.column_dimensions[get_column_letter(index)].width = \
            min(max(largeur + 2, 10), LARGEUR_MAX)


def _tableau(feuille, colonnes: List[str], lignes) -> None:
    _entete(feuille, colonnes)
    for ligne in lignes:
        feuille.append(ligne)
    if feuille.max_row > 1:
        feuille.auto_filter.ref = feuille.dimensions
    _ajuster(feuille)


def construire_classeur(modele, traduire: Callable[..., str],
                        contexte: Dict[str, str]) -> bytes:
    """Produit le classeur et rend ses octets."""
    t = lambda cle, params=None: traduire(cle, params or {})

    classeur = Workbook()

    # --- Feuille 1 : un rôle par ligne -----------------------------------
    feuille = classeur.active
    feuille.title = t("export.sheet.roles")[:31]
    _tableau(feuille, [
        t("export.column.role_id"), t("export.column.name"),
        t("export.column.description"), t("export.column.role_type"),
        t("export.column.status"), t("export.column.right_count"),
        t("export.column.member_count"), t("export.column.members_listed"),
        t("export.column.applications"),
        # Ce que le destinataire doit savoir avant d'intégrer : le rôle
        # est-il neuf, ou est-ce le même rôle dans une version différente ?
        # Sans cette colonne, il crée un doublon dans l'IGA.
        t("export.column.version"), t("export.column.change_since_export"),
    ], [
        [role.id, role.nom, role.description,
         t(f"export.role_type.{role.type_role}"),
         t(f"graph.origin.{role.origine}"),
         len(role.droits), role.membres_total, len(role.membres),
         ", ".join(role.applications),
         role.version,
         t(f"export.change.{role.evolution}") if role.evolution else ""]
        for role in modele.roles
    ])

    # --- Feuille 2 : un couple (rôle, droit) par ligne --------------------
    _tableau(classeur.create_sheet(t("export.sheet.rights")[:31]), [
        t("export.column.role_id"), t("export.column.name"),
        t("export.column.role_type"), t("export.column.status"),
        t("export.column.right_id"), t("export.column.application"),
    ], [
        [role.id, role.nom, t(f"export.role_type.{role.type_role}"),
         t(f"graph.origin.{role.origine}"), droit,
         modele_application(modele, droit)]
        for role in modele.roles for droit in role.droits
    ])

    # --- Feuille 3 : un couple (rôle, membre) par ligne -------------------
    _tableau(classeur.create_sheet(t("export.sheet.members")[:31]), [
        t("export.column.role_id"), t("export.column.name"),
        t("export.column.role_type"), t("export.column.status"),
        t("export.column.identity_id"),
    ], [
        [role.id, role.nom, t(f"export.role_type.{role.type_role}"),
         t(f"graph.origin.{role.origine}"), membre]
        for role in modele.roles for membre in role.membres
    ])

    # --- Feuille 4 : ce que cet extrait contient --------------------------
    _feuille_perimetre(classeur.create_sheet(t("export.sheet.scope")[:31]),
                       modele, t, contexte)

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def modele_application(modele, droit: str) -> str:
    """Application d'un droit, telle que les rôles la portent."""
    for role in modele.roles:
        if droit in role.droits and len(role.applications) == 1:
            return role.applications[0]
    return ""


def _feuille_perimetre(feuille, modele, t, contexte: Dict[str, str]) -> None:
    lignes: List[List[Any]] = [
        [t("export.scope.product"), NOM_PRODUIT],
        [t("export.scope.generated_at"),
         datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")],
        [t("export.scope.workspace"), contexte.get("workspace", "")],
        [t("export.scope.client"), contexte.get("client", "")],
        [t("export.scope.environment"), contexte.get("environnement", "")],
        [t("export.scope.author"), contexte.get("auteur", "")],
        [],
    ]

    for entree in modele.filtres.description():
        libelle = t(f"export.filter.{entree['critere']}")
        if entree["critere"] == "attribute":
            libelle = f"{libelle} · {entree['attribut']}"
        valeurs = entree["valeurs"]
        if entree["critere"] in ("origins", "role_types"):
            prefixe = ("graph.origin." if entree["critere"] == "origins"
                       else "export.role_type.")
            valeurs = [t(prefixe + str(v)) for v in valeurs]
        lignes.append([libelle, ", ".join(str(v) for v in valeurs)])

    lignes.append([])
    totaux = modele.totaux
    lignes.append([t("export.scope.total_roles"), totaux["roles"]])
    lignes.append([t("export.scope.total_rights"), totaux["droits"]])
    lignes.append([t("export.scope.total_identities"), totaux["identites"]])
    lignes.append([t("export.scope.last_export"),
                   contexte.get("dernier_export", "")
                   or t("export.scope.never_exported")])
    lignes.append([t("export.scope.new_roles"), totaux["nouveaux"]])
    lignes.append([t("export.scope.changed_roles"), totaux["modifies"]])

    masques = sum(role.membres_masques for role in modele.roles)
    if masques:
        # Le total d'un rôle et le nombre de lignes de la feuille des membres
        # divergent dès qu'un filtre identitaire s'applique. Non dit, cet écart
        # passerait pour une incohérence du produit.
        lignes.append([])
        lignes.append([t("export.scope.filtered_members_note"), masques])

    _entete(feuille, [t("export.column.criterion"), t("export.column.value")])
    for ligne in lignes:
        feuille.append(ligne)
    _ajuster(feuille)
