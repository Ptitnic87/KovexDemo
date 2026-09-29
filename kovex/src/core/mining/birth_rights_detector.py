# src/core/mining/birth_rights_detector.py

import pandas as pd
import numpy as np
from typing import Any, Callable, Dict, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
import os

from src.infrastructure import palette
from src.infrastructure.i18n_manager import i18n

#: Nom du rôle socle produit par la détection.
#:
#: Ce n'est pas un libellé mais un identifiant : il part dans les exports, et
#: de là dans l'outil de gouvernance du client, où il doit rester le même d'un
#: export à l'autre et d'une langue à l'autre. Il ne passe donc pas par le
#: catalogue de traductions — ce que le produit affiche à côté de lui, si, et
#: c'est la description.
#:
#: Il reste écrit ici plutôt que dans un réglage tant que personne n'a demandé
#: à le changer : un réglage que nul ne touche est une surface de plus à
#: tester. Le jour où un client impose sa nomenclature, c'est cette constante
#: qui devient un réglage, et rien d'autre ne bouge.
NOM_DU_ROLE_SOCLE = "ROLE_SOCLE_UNIVERSEL"

#: Nombre de droits portés par le diagramme de Pareto du classeur.
#:
#: Le chiffre était écrit trois fois : deux bornes de graphique exprimées en
#: lignes de tableau — donc en « 51 », en-tête comprise — et une troisième en
#: toutes lettres dans le titre de l'axe. Trois écritures d'une même valeur,
#: dont deux décalées d'une unité.
TAILLE_DU_PARETO = 50


class BirthRightsDetector:
    """
    Détecte les droits socles, qui polluent les analyses de mining.
    Ces droits apparaissent chez la majorité des utilisateurs et ne différencient pas les rôles.
    """
    
    def __init__(self, loader):
        """
        Args:
            loader: Instance du DataLoader contenant les habilitations
        """
        self.loader = loader
        self.habilitations = loader.habilitations
        self.identities = loader.identities
        self.rights = loader.rights
        
    def detect_birth_rights(self, frequency_threshold: float = 90.0) -> Dict[str, Any]:
        """
        Détecte les droits dont la fréquence dépasse le seuil (droits socles).
        
        Args:
            frequency_threshold: Pourcentage minimum (0-100) pour qu'un droit soit considéré comme un droit socle
            
        Returns:
            Dictionnaire contenant:
            - birth_rights: Liste des droits socles avec leurs stats
            - socle_role: Description du rôle socle créé
            - stats: Statistiques globales
            - pareto_data: Données pour le graphique de Pareto
        """
        
        if self.habilitations.empty or self.identities.empty:
            return self._empty_response()
        
        # 1. Calculer la fréquence de chaque droit
        #
        # Le dénominateur est la population que le mining regarde : les
        # identités qui détiennent au moins une habilitation. Le référentiel
        # entier comptait aussi celles qui n'en ont aucune — 3 573 sur 21 454,
        # soit 16,7 %, sur un référentiel réel — et plafonnait mécaniquement la
        # prévalence : un droit détenu par 100 % des identités actives
        # n'atteignait que 83,3 %, donc jamais le seuil de 90 % proposé par
        # défaut. La détection ne pouvait pas trouver ce qu'elle cherche.
        total_identities = len(self.identities)
        total_users = int(self.habilitations['ID_utilisateur'].nunique())
        
        # Compter combien d'utilisateurs ont chaque droit
        right_freq = self.habilitations.groupby('ID_droit')['ID_utilisateur'].nunique().reset_index()
        right_freq.columns = ['ID_droit', 'user_count']
        right_freq['frequency_pct'] = (right_freq['user_count'] / total_users) * 100
        
        # Trier par fréquence décroissante
        right_freq = right_freq.sort_values('frequency_pct', ascending=False)
        
        # 2. Identifier les droits socles (au-dessus du seuil)
        birth_rights_df = right_freq[right_freq['frequency_pct'] >= frequency_threshold].copy()
        
        # Enrichir avec les métadonnées des droits (si disponibles)
        if 'ID_droit' in self.rights.columns:
            birth_rights_df = birth_rights_df.merge(
                self.rights,
                on='ID_droit',
                how='left'
            )
        
        # 3. Calculer le Pareto (Cumul)
        right_freq['cumulative_pct'] = right_freq['frequency_pct'].cumsum() / right_freq['frequency_pct'].sum() * 100
        
        # 4. Préparer les résultats
        birth_rights_list = birth_rights_df.to_dict('records')
        
        # Nettoyer les valeurs NaN pour le JSON
        for br in birth_rights_list:
            for key, value in br.items():
                if pd.isna(value):
                    br[key] = None
        
        # 5. Créer le rôle socle
        socle_role = {
            "name": NOM_DU_ROLE_SOCLE,
            # La description n'est pas une phrase mais une clé et ses
            # paramètres : le serveur ne sait pas dans quelle langue elle sera
            # lue, et la composer ici la figeait en français jusque dans les
            # exports.
            "description_key": "birth_rights.socle_role.description",
            "description_params": {"rights": len(birth_rights_df),
                                   "threshold": frequency_threshold},
            "role_type": "SOCLE",
            "rights": birth_rights_df['ID_droit'].tolist(),
            "user_count": total_users,
            "right_count": len(birth_rights_df),
            "average_frequency": float(birth_rights_df['frequency_pct'].mean())
        }
        
        # 6. Statistiques globales
        stats = {
            "total_rights_analyzed": len(right_freq),
            "birth_rights_count": len(birth_rights_df),
            "birth_rights_percentage": round((len(birth_rights_df) / len(right_freq)) * 100, 1),
            "threshold_used": frequency_threshold,
            "total_users": total_users,
            # Ce que `total_users` compte, et ce qu'il ne compte pas : sans
            # cela, un pourcentage de prévalence n'est pas interprétable.
            "total_identities": total_identities,
            "identities_without_rights": total_identities - total_users,
            "avg_frequency_birth_rights": round(birth_rights_df['frequency_pct'].mean(), 1),
            "min_frequency_birth_rights": round(birth_rights_df['frequency_pct'].min(), 1) if len(birth_rights_df) > 0 else 0,
            "max_frequency_birth_rights": round(birth_rights_df['frequency_pct'].max(), 1) if len(birth_rights_df) > 0 else 0
        }
        
        # 7. Données pour le graphique de Pareto (Top 50 pour la lisibilité)
        pareto_data = right_freq.head(50)[['ID_droit', 'frequency_pct', 'cumulative_pct']].to_dict('records')
        
        # Nettoyer les NaN
        for item in pareto_data:
            for key, value in item.items():
                if pd.isna(value):
                    item[key] = None
        
        return {
            "birth_rights": birth_rights_list,
            "socle_role": socle_role,
            "stats": stats,
            "pareto_data": pareto_data
        }
    
    def export_to_excel(self, detection_result: Dict[str, Any],
                        output_path: str = "output/01_birth_rights_candidates.xlsx",
                        traduire: Optional[Callable[..., str]] = None) -> str:
        """Exporte les résultats vers un classeur Excel, dans la langue demandée.

        Ce classeur était le dernier livrable du produit écrit en français dans
        le code : titres de feuilles, en-têtes de colonnes, libellés du rôle
        socle, titres du graphique. Un consultant qui travaille en anglais
        recevait une interface anglaise et un classeur français.

        `traduire(cle, params)` est fourni par la route, qui connaît la langue
        demandée — c'est la convention des autres exports du produit. Sans lui,
        le classeur sort dans la langue de référence : un appel direct au
        moteur, hors requête, reste possible.

        Args:
            detection_result: Résultat de detect_birth_rights()
            output_path: Chemin du fichier Excel à créer
            traduire: Fonction de traduction `(clé, params) -> texte`

        Returns:
            Chemin du fichier créé
        """
        t = traduire or (lambda cle, params=None: i18n.t(cle, **(params or {})))
        
        # S'assurer que le dossier output existe
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Créer un workbook
        wb = openpyxl.Workbook()
        
        # --- FEUILLE 1 : Droits socles détectés ---
        ws_birth = wb.active
        ws_birth.title = t("birth_rights.export.sheet.rights")
        
        # En-têtes
        headers = [t("birth_rights.export.column.right_id"),
                   t("birth_rights.export.column.frequency"),
                   t("birth_rights.export.column.user_count"),
                   t("birth_rights.export.column.threshold")]
        ws_birth.append(headers)
        
        # Mise en forme des en-têtes
        # La palette du serveur, pas deux hexadécimaux recopiés. Elle existe
        # pour ça, et son palier de bandeau est celui qui tient le contraste
        # sur fond clair : le bleu employé ici était le palier 500, à 3,67:1
        # sous du blanc — le défaut même que l'interface a corrigé au lot 13.
        header_fill = PatternFill(start_color=palette.BANDEAU.lstrip("#"),
                                  end_color=palette.BANDEAU.lstrip("#"),
                                  fill_type="solid")
        header_font = Font(bold=True, color=palette.SUR_BANDEAU.lstrip("#"))
        
        for col_num, header in enumerate(headers, 1):
            cell = ws_birth.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        # Données
        threshold = detection_result['stats']['threshold_used']
        for br in detection_result['birth_rights']:
            ws_birth.append([
                br['ID_droit'],
                round(br['frequency_pct'], 2),
                br['user_count'],
                f"≥{threshold}%"
            ])
        
        # Ajuster les largeurs de colonnes
        ws_birth.column_dimensions['A'].width = 30
        ws_birth.column_dimensions['B'].width = 15
        ws_birth.column_dimensions['C'].width = 18
        ws_birth.column_dimensions['D'].width = 12
        
        # --- FEUILLE 2 : Données Pareto (pour le graphique) ---
        ws_pareto = wb.create_sheet(title=t("birth_rights.export.sheet.pareto"))
        
        pareto_headers = [t("birth_rights.export.column.right_id"),
                          t("birth_rights.export.column.frequency"),
                          t("birth_rights.export.column.cumulative")]
        ws_pareto.append(pareto_headers)
        
        # Mise en forme
        for col_num, header in enumerate(pareto_headers, 1):
            cell = ws_pareto.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        # Données
        for item in detection_result['pareto_data']:
            ws_pareto.append([
                item['ID_droit'],
                round(item['frequency_pct'], 2),
                round(item['cumulative_pct'], 2)
            ])
        
        # Ajuster les largeurs
        ws_pareto.column_dimensions['A'].width = 30
        ws_pareto.column_dimensions['B'].width = 15
        ws_pareto.column_dimensions['C'].width = 15
        
        # --- FEUILLE 3 : Rôle Socle ---
        ws_socle = wb.create_sheet(title=t("birth_rights.export.sheet.socle_role"))
        
        socle = detection_result['socle_role']
        
        ws_socle.append([t("birth_rights.export.column.property"),
                         t("birth_rights.export.column.value")])
        ws_socle.cell(row=1, column=1).fill = header_fill
        ws_socle.cell(row=1, column=1).font = header_font
        ws_socle.cell(row=1, column=2).fill = header_fill
        ws_socle.cell(row=1, column=2).font = header_font
        
        ws_socle.append([t("birth_rights.export.property.name"), socle['name']])
        ws_socle.append([t("birth_rights.export.property.description"),
                         t(socle['description_key'], socle['description_params'])])
        ws_socle.append([t("birth_rights.export.property.type"), socle['role_type']])
        ws_socle.append([t("birth_rights.export.property.right_count"), socle['right_count']])
        ws_socle.append([t("birth_rights.export.property.user_count"), socle['user_count']])
        ws_socle.append([t("birth_rights.export.property.average_frequency"),
                         f"{socle['average_frequency']:.1f}%"])
        
        ws_socle.column_dimensions['A'].width = 25
        ws_socle.column_dimensions['B'].width = 50
        
        # --- AJOUTER UN GRAPHIQUE PARETO ---
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = t("birth_rights.export.chart.title")
        chart.y_axis.title = t("birth_rights.export.column.frequency")
        chart.x_axis.title = t("birth_rights.export.chart.x_axis",
                               {"count": TAILLE_DU_PARETO})
        
        # Données : colonnes B (fréquence)
        data = Reference(ws_pareto, min_col=2, min_row=1, max_row=min(TAILLE_DU_PARETO + 1,
                                                 len(detection_result['pareto_data']) + 1))
        cats = Reference(ws_pareto, min_col=1, min_row=2, max_row=min(TAILLE_DU_PARETO + 1,
                                                 len(detection_result['pareto_data']) + 1))
        
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 12
        chart.width = 25
        
        ws_pareto.add_chart(chart, "E2")
        
        # Sauvegarder
        wb.save(output_path)
        
        return output_path
    
    def _empty_response(self) -> Dict[str, Any]:
        """Réponse vide si pas de données."""
        return {
            "birth_rights": [],
            "socle_role": {
                "name": NOM_DU_ROLE_SOCLE,
                "description_key": "birth_rights.socle_role.empty",
                "description_params": {},
                "role_type": "SOCLE",
                "rights": [],
                "user_count": 0,
                "right_count": 0,
                "average_frequency": 0
            },
            "stats": {
                "total_rights_analyzed": 0,
                "birth_rights_count": 0,
                "birth_rights_percentage": 0,
                "threshold_used": 90.0,
                "total_users": 0,
                "total_identities": 0,
                "identities_without_rights": 0,
                "avg_frequency_birth_rights": 0,
                "min_frequency_birth_rights": 0,
                "max_frequency_birth_rights": 0
            },
            "pareto_data": []
        }
